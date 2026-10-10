"""Genera el libro de respuestas de la Encuesta de Cultura HSE 2026.

Uso: python plantilla.py headers.json salida.xlsx [--ejemplo]
  --ejemplo agrega filas de prueba (solo para verificar las fórmulas, no se entrega).
"""
import json
import random
import sys
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

HEAD = json.load(open(sys.argv[1], encoding="utf-8"))   # las 46 columnas del Form, texto exacto
SALIDA = sys.argv[2]
EJEMPLO = "--ejemplo" in sys.argv
assert len(HEAD) == 46

HOJA = "Respuestas de formulario 1"
REF = "'" + HOJA + "'"
FILAS = 5000            # alcance de las fórmulas del resumen
F = "Arial"

# Mismos textos que index.html (OPCIONES y ESCALA): es lo que se guarda en la hoja.
OPC = {
    2: ["AGG Planta Mondoñedo", "CEM Planta Nobsa", "CORPORATIVO Teleport", "CORPORATIVO Fundación Holcim",
        "GEOCYCLE Barrancabermeja", "GEOCYCLE Nobsa", "RMX Planta Bello", "RMX Plantas Valle", "RMX Planta Chía",
        "RMX Plantas Boyaca", "RMX Planta Puente Aranda", "RMX Planta Sibaté", "TQC Tocancipá"],
    3: ["Empleado(a) Holcim", "Contratista Holcim"],
    4: ["Director(a)", "Gerente - Head", "Jefe - Coordinador(a)", "Supervisor(a) - Analista", "Operativo(a)"],
    5: ["Menos de 20", "20 - 30", "31 - 40", "41 - 50", "51 - 60", "+ 61"],
    6: ["0 a 1", "1 a 5", "6 a 10", "11 a 15", "16 a 20", "+ 20"],
    7: ["Si", "No"],
}
ESCALA = ["Totalmente de Acuerdo", "De acuerdo", "Un poco de Acuerdo", "Ni de acuerdo ni en desacuerdo",
          "Un poco en Desacuerdo", "En Desacuerdo", "Totalmente en Desacuerdo"]
FAVORABLE, NEUTRAL, DESFAVORABLE = ESCALA[:3], ESCALA[3:4], ESCALA[4:]

# Dimensiones: (nombre, primera y última columna de afirmaciones, columna de comentarios)
DIM = [
    ("Valores y compromiso HSE", 8, 11, 12),
    ("Procesos operacionales", 13, 16, 17),
    ("Gestión de recursos", 18, 21, 22),
    ("Supervisión", 23, 26, 27),
    ("Planificación del trabajo", 28, 31, 32),
    ("Abordaje de problemas y peligros", 33, 36, 37),
    ("Cumplimiento de reglas", 38, 41, 42),
    ("Evaluación subjetiva de riesgos", 43, 45, 46),
]

NAVY, VERDE, CIAN, GRIS, BLANCO = "071A38", "97D13B", "1BB4E4", "EEF3F9", "FFFFFF"
fino = Side(style="thin", color="D7E0EC")
borde = Border(left=fino, right=fino, top=fino, bottom=fino)


def cab(c, fondo=NAVY, color=BLANCO):
    c.font = Font(name=F, bold=True, color=color, size=10)
    c.fill = PatternFill("solid", fgColor=fondo)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = borde


def seccion_de(i):
    if i == 1:
        return "Registro"
    if i <= 7:
        return "Datos generales"
    for n, a, b, com in DIM:
        if a <= i <= com:
            return n
    return ""


wb = Workbook()

# ---------------------------------------------------------------- Respuestas
ws = wb.active
ws.title = HOJA
for i, h in enumerate(HEAD, 1):
    c = ws.cell(row=1, column=i, value=h)
    # Datos generales en azul, afirmaciones en navy, comentarios en gris
    if 2 <= i <= 7:
        cab(c, "1877C4")
    elif any(com == i for _, _, _, com in DIM):
        cab(c, "5B6B82")
    else:
        cab(c)
    ws.column_dimensions[L(i)].width = 22 if i == 1 else (34 if any(com == i for *_, com in DIM) else 26)
ws.row_dimensions[1].height = 120
ws.freeze_panes = "B2"
ws.auto_filter.ref = "A1:%s1" % L(46)
for fila in ws.iter_rows(min_row=2, max_row=200, max_col=46):
    for c in fila:
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(vertical="top", wrap_text=c.column in [com for *_, com in DIM])
for r in range(2, FILAS + 1):
    ws.cell(row=r, column=1).number_format = "dd/mm/yyyy hh:mm:ss"
ws["A1"].comment = Comment("La escribe sola la encuesta (fecha y hora de Bogotá). No cambie los títulos de esta fila: "
                           "el script ubica cada respuesta por el texto del título.", "Encuesta")

if EJEMPLO:
    random.seed(7)
    base = datetime(2026, 10, 1, 7, 0)
    for r in range(2, 22):
        fila = [base + timedelta(hours=r * 5)]
        for i in range(2, 47):
            if i in OPC:
                fila.append(random.choice(OPC[i]))
            elif any(com == i for *_, com in DIM):
                fila.append(random.choice(["", "", "Buen trabajo en campo"]))
            else:
                fila.append(random.choice(ESCALA))
        for i, v in enumerate(fila, 1):
            ws.cell(row=r, column=i, value=v)

# ---------------------------------------------------------------- Reportes
rp = wb.create_sheet("Reportes de problemas")
ENC_REP = ["Marca temporal", "Tipo", "Descripción", "Pantalla", "Contacto", "Navegador",
           "Tamaño de pantalla", "Tema", "Último error", "Estado"]
ANCHO_REP = [20, 24, 50, 30, 26, 40, 16, 10, 36, 12]
for i, (h, w) in enumerate(zip(ENC_REP, ANCHO_REP), 1):
    cab(rp.cell(row=1, column=i, value=h))
    rp.column_dimensions[L(i)].width = w
rp.freeze_panes = "A2"
rp.auto_filter.ref = "A1:J1"
for r in range(2, 1001):
    rp.cell(row=r, column=1).number_format = "dd/mm/yyyy hh:mm:ss"
rp["J1"].comment = Comment("Llega en 'Nuevo'. Cámbielo a 'En revisión' o 'Resuelto' para hacer seguimiento.", "Encuesta")

# ---------------------------------------------------------------- Resumen
rs = wb.create_sheet("Resumen", 0)
rs.sheet_view.showGridLines = False
rs.column_dimensions["A"].width = 38
for col in "BCDEF":
    rs.column_dimensions[col].width = 16
rs["A1"] = "Encuesta de Cultura de Salud, Seguridad y Ambiente 2026 — Resumen"
rs["A1"].font = Font(name=F, bold=True, size=15, color=NAVY)
rs["A2"] = "Se calcula solo con lo que llega a la hoja '" + HOJA + "'. No escriba en esta hoja."
rs["A2"].font = Font(name=F, italic=True, size=9, color="5B6B82")

def etiqueta(celda, texto, negrita=False):
    rs[celda] = texto
    rs[celda].font = Font(name=F, bold=negrita, size=10)

A_RANGO = "%s!$A$2:$A$%d" % (REF, FILAS)
etiqueta("A4", "Respuestas recibidas", True)
rs["B4"] = "=COUNTA(%s)" % A_RANGO
etiqueta("A5", "Primera respuesta", True)
rs["B5"] = '=IF(B4=0,"",MIN(%s))' % A_RANGO
etiqueta("A6", "Última respuesta", True)
rs["B6"] = '=IF(B4=0,"",MAX(%s))' % A_RANGO
for c in ("B4", "B5", "B6"):
    rs[c].font = Font(name=F, bold=True, size=11, color=NAVY)
    rs[c].alignment = Alignment(horizontal="left")
rs["B5"].number_format = rs["B6"].number_format = "dd/mm/yyyy hh:mm"

fila = 8

def tabla_conteo(titulo, col, opciones):
    global fila
    rs.cell(row=fila, column=1, value=titulo)
    cab(rs.cell(row=fila, column=1), NAVY)
    rs.cell(row=fila, column=1).alignment = Alignment(horizontal="left", vertical="center")
    for j, t in enumerate(["Respuestas", "% del total"], 2):
        cab(rs.cell(row=fila, column=j, value=t))
    fila += 1
    rango = "%s!$%s$2:$%s$%d" % (REF, L(col), L(col), FILAS)
    for o in opciones:
        rs.cell(row=fila, column=1, value=o).font = Font(name=F, size=10)
        rs.cell(row=fila, column=2, value='=COUNTIF(%s,A%d)' % (rango, fila)).font = Font(name=F, size=10)
        c = rs.cell(row=fila, column=3, value='=IF($B$4=0,"",B%d/$B$4)' % fila)
        c.number_format = "0.0%"; c.font = Font(name=F, size=10)
        for j in (1, 2, 3):
            rs.cell(row=fila, column=j).border = borde
        fila += 1
    fila += 1

tabla_conteo("Planta o sede", 2, OPC[2])
tabla_conteo("Vinculación", 3, OPC[3])
tabla_conteo("Tipo de cargo", 4, OPC[4])
tabla_conteo("¿Ha reportado desviaciones sin que se tengan en cuenta?", 7, OPC[7])

# % favorable / neutral / desfavorable por dimensión
rs.cell(row=fila, column=1, value="Percepción por dimensión")
for j, t in enumerate(["Calificaciones", "% Favorable", "% Neutral", "% Desfavorable"], 2):
    cab(rs.cell(row=fila, column=j, value=t))
cab(rs.cell(row=fila, column=1)); rs.cell(row=fila, column=1).alignment = Alignment(horizontal="left", vertical="center")
nota = fila
fila += 1
for n, a, b, _ in DIM:
    rango = "%s!$%s$2:$%s$%d" % (REF, L(a), L(b), FILAS)
    suma = lambda ops: "+".join('COUNTIF(%s,"%s")' % (rango, o) for o in ops)
    rs.cell(row=fila, column=1, value=n).font = Font(name=F, size=10)
    rs.cell(row=fila, column=2, value="=COUNTA(%s)" % rango).font = Font(name=F, size=10)
    for j, ops in ((3, FAVORABLE), (4, NEUTRAL), (5, DESFAVORABLE)):
        c = rs.cell(row=fila, column=j, value='=IF($B%d=0,"",(%s)/$B%d)' % (fila, suma(ops), fila))
        c.number_format = "0.0%"; c.font = Font(name=F, size=10)
    for j in range(1, 6):
        rs.cell(row=fila, column=j).border = borde
    fila += 1
rs.cell(row=nota, column=3).comment = Comment(
    "Favorable = Totalmente de Acuerdo + De acuerdo + Un poco de Acuerdo. Neutral = Ni de acuerdo ni en desacuerdo. "
    "Desfavorable = las tres de desacuerdo. Se cuenta sobre todas las afirmaciones de la dimensión.", "Encuesta")
fila += 1
rs.cell(row=fila, column=1, value="Favorable = las tres opciones de acuerdo · Neutral = Ni de acuerdo ni en desacuerdo · "
        "Desfavorable = las tres opciones de desacuerdo.").font = Font(name=F, italic=True, size=9, color="5B6B82")
rs.freeze_panes = "A4"

# ---------------------------------------------------------------- Diccionario
dc = wb.create_sheet("Diccionario")
for i, (h, w) in enumerate(zip(["Col.", "N.º", "Sección", "Pregunta (título exacto de la columna)", "Tipo", "Valores que puede tener"],
                               [7, 6, 30, 70, 18, 70]), 1):
    cab(dc.cell(row=1, column=i, value=h)); dc.column_dimensions[L(i)].width = w
dc.freeze_panes = "A2"
for i, h in enumerate(HEAD, 1):
    if i == 1:
        tipo, val = "Fecha y hora", "La escribe sola la encuesta"
    elif i in OPC:
        tipo, val = "Opción", " · ".join(OPC[i])
    elif any(com == i for *_, com in DIM):
        tipo, val = "Texto (opcional)", "Comentario libre, hasta 1000 caracteres"
    else:
        tipo, val = "Escala de acuerdo", " · ".join(ESCALA)
    for j, v in enumerate([L(i), i, seccion_de(i), h.strip(), tipo, val], 1):
        c = dc.cell(row=i + 1, column=j, value=v)
        c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = borde

# ---------------------------------------------------------------- Léame
lm = wb.create_sheet("Léame", 0)
lm.sheet_view.showGridLines = False
lm.column_dimensions["A"].width = 4
lm.column_dimensions["B"].width = 110
textos = [
    ("Encuesta de Cultura de Salud, Seguridad y Ambiente 2026 — libro de respuestas", "titulo"),
    ("Cada respuesta de la encuesta (https://zuicagerman-eng.github.io/Encuesta-Cultura-y-seguradad/) llega como una fila nueva, igual que en Google Forms.", ""),
    ("", ""),
    ("Hojas", "sub"),
    ("Respuestas de formulario 1 — una fila por persona, con las 46 columnas del Form en el mismo orden. La escribe la encuesta: no la edite a mano.", ""),
    ("Reportes de problemas — lo que la gente envía desde el botón Ayuda. Use la columna Estado para hacer seguimiento.", ""),
    ("Resumen — total de respuestas, respuestas por planta, vinculación y cargo, y % favorable por dimensión. Se actualiza solo.", ""),
    ("Diccionario — qué es cada columna y qué valores puede tener.", ""),
    ("", ""),
    ("Cómo conectarlo", "sub"),
    ("1. Suba este archivo a Google Drive y ábralo con Google Sheets (Archivo → Guardar como Hojas de cálculo de Google).", ""),
    ("2. En la hoja de Google: Extensiones → Apps Script. Pegue apps-script/Codigo.gs del repositorio y guarde.", ""),
    ("3. Ejecute la función prepararHoja: debe decir que están las 46 columnas.", ""),
    ("4. Implementar → Nueva implementación → Aplicación web. Ejecutar como: Yo. Acceso: Cualquier persona. Copie la URL /exec.", ""),
    ("5. Esa URL va en URL_WEBAPP dentro de index.html.", ""),
    ("", ""),
    ("No cambie los títulos de la fila 1 de 'Respuestas de formulario 1': el script ubica cada respuesta por ese texto. "
     "Si lo cambia, la encuesta avisa con un error en vez de guardar en la columna equivocada.", "ojo"),
]
for k, (t, estilo) in enumerate(textos, 2):
    c = lm.cell(row=k, column=2, value=t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if estilo == "titulo":
        c.font = Font(name=F, bold=True, size=15, color=NAVY)
    elif estilo == "sub":
        c.font = Font(name=F, bold=True, size=12, color="1877C4")
    elif estilo == "ojo":
        c.font = Font(name=F, bold=True, size=10, color="A12A22")
        c.fill = PatternFill("solid", fgColor="FDE7E5")
    else:
        c.font = Font(name=F, size=10)

wb.active = 0
wb.save(SALIDA)
print("guardado", SALIDA)
